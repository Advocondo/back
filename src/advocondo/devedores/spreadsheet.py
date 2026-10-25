"""The spreadsheet of residents in default sent by the condominium (US14)."""

import re
from dataclasses import dataclass
from enum import StrEnum
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from advocondo.common.normalization import comparable_text
from advocondo.devedores.errors import InvalidSpreadsheetError
from advocondo.devedores.models import MAX_NAME_LENGTH, MAX_UNIT_LENGTH

MAX_FILE_BYTES = 1024 * 1024
MAX_ROWS = 1000
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ImportRowStatus(StrEnum):
    VALID = "valid"
    INVALID = "invalid"
    # Looks like a registered debtor or an earlier row: skipped, as in US13-CA01.
    SIMILAR = "similar"


@dataclass(frozen=True)
class Column:
    header: str
    field: str
    required: bool
    # Shown on the row when the value is refused, whatever the reason.
    invalid_message: str


COLUMNS = (
    Column(
        "Nome", "name", True, f"Informe o nome, com até {MAX_NAME_LENGTH} caracteres."
    ),
    Column(
        "Unidade",
        "unit",
        True,
        f"Informe a unidade, com até {MAX_UNIT_LENGTH} caracteres.",
    ),
    Column(
        "Telefone", "phone", False, "Telefone inválido: informe DDD e 10 ou 11 dígitos."
    ),
    Column("E-mail", "email", False, "E-mail inválido."),
)


@dataclass(frozen=True)
class SpreadsheetRow:
    number: int
    """Row number as Excel shows it, so the user can find the row to fix."""
    values: dict[str, str | None]
    """Cell text by request field (`name`, `unit`, `phone`, `email`)."""


def read_rows(content: bytes) -> list[SpreadsheetRow]:
    """Rows of the first sheet, read by the headers in the first row."""
    if len(content) > MAX_FILE_BYTES:
        raise InvalidSpreadsheetError(
            f"A planilha passa de {MAX_FILE_BYTES // 1024 // 1024} MB."
        )
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, ValueError) as error:
        raise InvalidSpreadsheetError() from error
    try:
        lines = workbook.worksheets[0].iter_rows(values_only=True)
        fields_by_position = _fields_by_position(next(lines, ()))
        rows: list[SpreadsheetRow] = []
        for number, line in enumerate(lines, start=2):
            values = {
                field: _cell_text(line[position]) if position < len(line) else None
                for position, field in fields_by_position.items()
            }
            if not any(values.values()):
                continue
            if len(rows) == MAX_ROWS:
                raise InvalidSpreadsheetError(
                    f"A planilha passa de {MAX_ROWS} linhas. Divida o arquivo."
                )
            rows.append(SpreadsheetRow(number, values))
    finally:
        workbook.close()
    if not rows:
        raise InvalidSpreadsheetError("A planilha não tem nenhum devedor.")
    return rows


def build_template() -> bytes:
    """Empty spreadsheet with the headers that `read_rows` expects."""
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Inadimplentes"
    worksheet.append([column.header for column in COLUMNS])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _header_key(header: object) -> str:
    """'E-mail', 'EMAIL' and ' e mail ' are the same header."""
    return re.sub(r"[^0-9a-z]", "", comparable_text(str(header or "")))


def _fields_by_position(headers: tuple[object, ...]) -> dict[int, str]:
    positions = {
        _header_key(header): position for position, header in enumerate(headers)
    }
    missing = [
        column.header
        for column in COLUMNS
        if column.required and _header_key(column.header) not in positions
    ]
    if missing:
        raise InvalidSpreadsheetError(
            "Faltam colunas na planilha: "
            + ", ".join(missing)
            + ". Use o modelo de planilha."
        )
    return {
        positions[_header_key(column.header)]: column.field
        for column in COLUMNS
        if _header_key(column.header) in positions
    }


def _cell_text(value: object) -> str | None:
    """A phone typed without mask comes as a number: 61912345678 -> '61912345678'."""
    text = str(value).strip() if value is not None else ""
    return text or None
