import uuid
from collections.abc import Callable
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from advocondo.devedores.errors import InvalidSpreadsheetError
from advocondo.devedores.models import Debtor
from advocondo.devedores.schemas import DebtorCreateRequest, DebtorUpdateRequest
from advocondo.devedores.service import check_row
from advocondo.devedores.similarity import comparable_unit, find_similar_debtors
from advocondo.devedores.spreadsheet import (
    MAX_FILE_BYTES,
    MAX_ROWS,
    ImportRowStatus,
    SpreadsheetRow,
    build_template,
    read_rows,
)
from tests.support.spreadsheet import HEADERS, build_xlsx

CONDOMINIUM_ID = uuid.uuid4()


def make_debtor(name: str, unit: str) -> Debtor:
    return Debtor(
        condominium_id=CONDOMINIUM_ID, name=name, unit=unit, phone="61912345678"
    )


class TestFindSimilarDebtors:
    @pytest.mark.parametrize(
        ("name", "unit", "similar"),
        [
            pytest.param("José Simão", "Torre A 101", True, id="same-name"),
            pytest.param(
                "JOSE  SIMAO", "Torre A 101", True, id="name-without-accents-or-case"
            ),
            pytest.param("Maria Souza", "Bloco B 304", True, id="same-unit"),
            pytest.param(
                "Maria Souza", "bloco b-304", True, id="unit-written-differently"
            ),
            pytest.param("Maria Souza", "Bloco C 12", False, id="other-name-and-unit"),
        ],
    )
    def test_debtor_with_same_name_or_unit_is_similar(
        self, name: str, unit: str, similar: bool
    ) -> None:
        """US13-CA01"""
        registered = make_debtor("José Simão", "Bloco B 304")

        expected = [registered] if similar else []
        assert find_similar_debtors(name, unit, [registered]) == expected


@pytest.mark.parametrize(
    ("typed", "same_as"),
    [
        pytest.param("Bloco B - 304", "blocob304", id="spaces-and-dash"),
        pytest.param("Torre A/101", "torre a 101", id="slash-and-case"),
        pytest.param("Apto. 12", "APTO12", id="dot"),
    ],
)
def test_comparable_unit_keeps_only_letters_and_numbers(
    typed: str, same_as: str
) -> None:
    assert comparable_unit(typed) == comparable_unit(same_as)


class TestDebtorCreateRequest:
    @pytest.mark.parametrize(
        ("typed", "stored"),
        [
            pytest.param("(61) 91234-5678", "61912345678", id="mobile-with-mask"),
            pytest.param("61 91234 5678", "61912345678", id="mobile-with-spaces"),
            pytest.param("(61) 3234-5678", "6132345678", id="landline-with-mask"),
            pytest.param("61912345678", "61912345678", id="digits-only"),
        ],
    )
    def test_masked_phone_is_stored_as_digits(self, typed: str, stored: str) -> None:
        request = DebtorCreateRequest(
            condominium_id=CONDOMINIUM_ID,
            name=" José Simão ",
            unit="Bloco B 304",
            phone=typed,
        )

        assert request.name == "José Simão"
        assert request.phone == stored
        assert request.confirm_duplicate is False

    @pytest.mark.parametrize(
        "fields",
        [
            pytest.param({"unit": "101"}, id="missing-name"),
            pytest.param({"name": "José"}, id="missing-unit"),
            pytest.param(
                {"name": "José", "unit": "101", "phone": "123"}, id="short-phone"
            ),
            pytest.param(
                {"name": "José", "unit": "101", "email": "jose@"}, id="invalid-email"
            ),
            pytest.param(
                {"name": "José", "unit": "101", "cpf": "1"}, id="unknown-field"
            ),
        ],
    )
    def test_invalid_request_is_refused(self, fields: dict[str, object]) -> None:
        with pytest.raises(ValidationError):
            DebtorCreateRequest.model_validate(
                {"condominium_id": str(CONDOMINIUM_ID), **fields}
            )


class TestDebtorUpdateRequest:
    @pytest.mark.parametrize(
        "fields",
        [
            pytest.param({}, id="nothing-to-change"),
            pytest.param({"name": None}, id="name-removed"),
            pytest.param({"unit": None}, id="unit-removed"),
        ],
    )
    def test_invalid_update_is_refused(self, fields: dict[str, object]) -> None:
        with pytest.raises(ValidationError):
            DebtorUpdateRequest.model_validate(fields)

    def test_new_name_and_unit_are_accepted(self) -> None:
        request = DebtorUpdateRequest.model_validate(
            {"name": "José Simão Filho", "unit": "Bloco C 101"}
        )

        assert (request.name, request.unit) == ("José Simão Filho", "Bloco C 101")


class TestReadSpreadsheet:
    """US14"""

    def test_rows_are_read_by_header_with_excel_row_numbers(self) -> None:
        content = build_xlsx(
            ("José Simão", "Bloco B 304", "(61) 91234-5678", None),
            (None, None, None, None),
            ("Maria Souza", "Torre A 101", 61987654321, " maria@exemplo.com "),
        )

        assert read_rows(content) == [
            SpreadsheetRow(
                2,
                {
                    "name": "José Simão",
                    "unit": "Bloco B 304",
                    "phone": "(61) 91234-5678",
                    "email": None,
                },
            ),
            SpreadsheetRow(
                4,
                {
                    "name": "Maria Souza",
                    "unit": "Torre A 101",
                    "phone": "61987654321",
                    "email": "maria@exemplo.com",
                },
            ),
        ]

    def test_headers_ignore_case_accents_order_and_punctuation(self) -> None:
        content = build_xlsx(
            ("maria@exemplo.com", "Torre A 101", "Maria"),
            headers=(" EMAIL ", "unidade", "NOME"),
        )

        assert read_rows(content)[0].values == {
            "name": "Maria",
            "unit": "Torre A 101",
            "email": "maria@exemplo.com",
        }

    @pytest.mark.parametrize(
        ("build", "message"),
        [
            pytest.param(
                lambda: build_xlsx(("José",), headers=("Nome", "Telefone")),
                "Faltam colunas.*Unidade",
                id="missing-required-column",
            ),
            pytest.param(build_xlsx, "nenhum devedor", id="no-debtors"),
            pytest.param(
                lambda: build_xlsx(
                    *[
                        (f"Devedor {n}", str(n), "61912345678")
                        for n in range(MAX_ROWS + 1)
                    ]
                ),
                "linhas",
                id="too-many-rows",
            ),
        ],
    )
    def test_spreadsheet_that_cannot_be_imported_is_refused(
        self, build: Callable[[], bytes], message: str
    ) -> None:
        """US14-CA01: o erro aparece antes de qualquer linha ser importada."""
        with pytest.raises(InvalidSpreadsheetError, match=message):
            read_rows(build())

    @pytest.mark.parametrize(
        "content",
        [
            pytest.param(b"Nome;Unidade", id="not-xlsx"),
            pytest.param(b"x" * (MAX_FILE_BYTES + 1), id="too-large"),
        ],
    )
    def test_unreadable_file_is_refused(self, content: bytes) -> None:
        with pytest.raises(InvalidSpreadsheetError):
            read_rows(content)

    def test_template_has_the_headers_read_by_the_import(self) -> None:
        worksheet = load_workbook(BytesIO(build_template())).worksheets[0]

        assert next(worksheet.iter_rows(values_only=True)) == HEADERS


def spreadsheet_row(
    values: dict[str, str | None] | None = None, number: int = 2
) -> SpreadsheetRow:
    defaults: dict[str, str | None] = {
        "name": "José Simão",
        "unit": "Bloco B 304",
        "phone": "(61) 91234-5678",
        "email": None,
    }
    return SpreadsheetRow(number, {**defaults, **(values or {})})


class TestCheckRow:
    """US14-CA01"""

    def test_valid_row_becomes_a_debtor_with_normalized_phone(self) -> None:
        row, debtor = check_row(CONDOMINIUM_ID, spreadsheet_row(), [], {})

        assert row.status is ImportRowStatus.VALID
        assert debtor is not None
        assert debtor.phone == "61912345678"
        assert debtor.condominium_id == CONDOMINIUM_ID

    @pytest.mark.parametrize(
        ("values", "expected_errors"),
        [
            pytest.param(
                {"name": None, "phone": "123"},
                [
                    "Informe o nome, com até 200 caracteres.",
                    "Telefone inválido: informe DDD e 10 ou 11 dígitos.",
                ],
                id="no-name-and-bad-phone",
            ),
            pytest.param({"email": "sem-arroba"}, ["E-mail inválido."], id="bad-email"),
            pytest.param(
                {"phone": None},
                ["Informe o telefone ou o e-mail do devedor."],
                id="no-contact",
            ),
        ],
    )
    def test_invalid_row_lists_its_errors(
        self, values: dict[str, str | None], expected_errors: list[str]
    ) -> None:
        row, debtor = check_row(CONDOMINIUM_ID, spreadsheet_row(values), [], {})

        assert row.status is ImportRowStatus.INVALID
        assert row.errors == expected_errors
        assert debtor is None

    def test_row_similar_to_a_registered_debtor_is_skipped(self) -> None:
        registered = make_debtor("Jose Simao", "Torre A 101")

        row, debtor = check_row(CONDOMINIUM_ID, spreadsheet_row(), [registered], {})

        assert row.status is ImportRowStatus.SIMILAR
        assert row.similar_debtors == [registered]
        assert debtor is None

    def test_row_repeating_an_earlier_row_is_skipped(self) -> None:
        earlier = make_debtor("Maria Souza", "Bloco B 304")

        row, debtor = check_row(
            CONDOMINIUM_ID, spreadsheet_row(number=5), [], {2: earlier}
        )

        assert row.status is ImportRowStatus.SIMILAR
        assert row.repeats_row == 2
        assert debtor is None
