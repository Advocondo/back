from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from advocondo.parcelas.status import InstallmentStatus

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = (
    "Condomínio",
    "Devedor",
    "Unidade",
    "ID do acordo",
    "Parcela",
    "Vencimento",
    "Pagamento",
    "Status",
    "Valor total",
    "Honorários",
    "Responsável",
)


@dataclass(frozen=True)
class ExportableInstallmentRow:
    condominium_name: str
    debtor_name: str
    debtor_unit: str
    agreement_id: str
    installment_number: int
    due_date: date
    payment_date: date | None
    status: InstallmentStatus
    amount: Decimal
    fee_amount: Decimal
    responsible_lawyer_name: str | None


def build_installment_workbook(rows: Sequence[ExportableInstallmentRow]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.create_sheet("Acompanhamento")
    default_sheet = workbook["Sheet"]
    workbook.remove(default_sheet)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:K{len(rows) + 1}"

    worksheet.append(HEADERS)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    for row_index, row in enumerate(rows, start=2):
        worksheet.append(
            (
                row.condominium_name,
                row.debtor_name,
                row.debtor_unit,
                row.agreement_id,
                row.installment_number,
                row.due_date,
                row.payment_date,
                row.status.value,
                row.amount,
                row.fee_amount,
                row.responsible_lawyer_name,
            )
        )
        # User-provided names and identifiers must remain text, never formulas.
        for column_index in (1, 2, 3, 4, 8, 11):
            worksheet.cell(row=row_index, column=column_index).data_type = "s"

    for excel_row in worksheet.iter_rows(min_row=2):
        for cell in (excel_row[5], excel_row[6]):
            cell.number_format = "dd/mm/yyyy"
        for cell in (excel_row[8], excel_row[9]):
            cell.number_format = "#,##0.00"

    for column_index, header in enumerate(HEADERS, start=1):
        cells = worksheet.iter_cols(
            min_col=column_index, max_col=column_index, min_row=1
        )
        values = next(cells)
        width = max(len(header), *(len(str(cell.value or "")) for cell in values))
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            width + 2, 40
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
