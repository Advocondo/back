from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from advocondo.common.errors import PermissionDeniedError
from advocondo.painel_admin.export import (
    HEADERS,
    ExportableInstallmentRow,
    build_installment_workbook,
)
from advocondo.painel_admin.repository import AdminPanelRepository
from advocondo.painel_admin.schemas import month_bounds
from advocondo.painel_admin.service import AdminPanelService
from advocondo.painel_operacional.schemas import (
    month_bounds as operational_month_bounds,
)
from advocondo.parcelas.status import InstallmentStatus
from advocondo.users.models import Role, User
from tests.support.clock import FrozenClock


@pytest.mark.parametrize("bounds", [month_bounds, operational_month_bounds])
@pytest.mark.parametrize(
    ("today", "expected"),
    [
        pytest.param(
            date(2026, 3, 2), (date(2026, 3, 1), date(2026, 4, 1)), id="march"
        ),
        pytest.param(
            date(2026, 12, 31), (date(2026, 12, 1), date(2027, 1, 1)), id="december"
        ),
    ],
)
def test_month_bounds_end_on_the_first_day_of_the_next_month(
    bounds: object, today: date, expected: tuple[date, date]
) -> None:
    """O fim é exclusivo, então dezembro termina em 1º de janeiro."""
    assert callable(bounds)
    assert bounds(today) == expected


def make_row(**fields: object) -> ExportableInstallmentRow:
    values: dict[str, object] = {
        "condominium_name": "Residencial Mirante",
        "debtor_name": "José",
        "debtor_unit": "Bloco B 304",
        "agreement_id": "a1",
        "installment_number": 1,
        "due_date": date(2026, 3, 10),
        "payment_date": None,
        "status": InstallmentStatus.OPEN,
        "amount": Decimal("1234.50"),
        "fee_amount": Decimal("123.45"),
        "responsible_lawyer_name": None,
    }
    values.update(fields)
    return ExportableInstallmentRow(**values)  # type: ignore[arg-type]


class TestInstallmentWorkbook:
    def test_empty_report_keeps_the_header(self) -> None:
        sheet = load_workbook(BytesIO(build_installment_workbook([])))["Acompanhamento"]

        assert [cell.value for cell in sheet[1]] == list(HEADERS)
        assert sheet.max_row == 1

    def test_money_is_a_number_with_two_decimals(self) -> None:
        """US35-CA02"""
        sheet = load_workbook(BytesIO(build_installment_workbook([make_row()])))[
            "Acompanhamento"
        ]

        amount, fee = sheet["I2"], sheet["J2"]
        assert amount.value == pytest.approx(1234.5)
        assert amount.number_format == "#,##0.00"
        assert fee.value == pytest.approx(123.45)

    def test_installment_without_lawyer_leaves_the_cell_empty(self) -> None:
        sheet = load_workbook(BytesIO(build_installment_workbook([make_row()])))[
            "Acompanhamento"
        ]

        assert sheet["K2"].value is None


class TestAdminPanelServicePermissions:
    """O router já exige Administrador; o service confere de novo, para quem o
    chamar de outro lugar."""

    @pytest.mark.parametrize("role", [Role.OPERATOR, Role.CLIENT])
    def test_only_the_admin_gets_the_panel_data(self, role: Role) -> None:
        service = AdminPanelService(
            AdminPanelRepository(session=None),  # type: ignore[arg-type]
            FrozenClock(),
        )
        user = User(role=role)

        with pytest.raises(PermissionDeniedError):
            service.get_financial_summary(user)
        with pytest.raises(PermissionDeniedError):
            service.get_dashboard(user, 2026, 3)
        with pytest.raises(PermissionDeniedError):
            service.export_monthly_followup(user, 2026, 3)
