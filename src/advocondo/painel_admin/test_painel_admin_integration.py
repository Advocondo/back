from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from advocondo.condominios.models import Condominium
from advocondo.users.models import Role, User
from tests.fixtures.auth import AuthenticateAs
from tests.fixtures.condominios import CreateCondominium
from tests.fixtures.devedores import CreateDebtor
from tests.fixtures.painel_cliente import CreateAgreementWithInstallments
from tests.fixtures.users import CreateUser

FINANCIAL_URL = "/admin-panel/financial"
DASHBOARD_URL = "/admin-panel/dashboard"
EXPORT_URL = "/admin-panel/reports/export"
# O relógio dos testes está em 2 de março de 2026: o mês atual é março.


@pytest.fixture
def admin_headers(
    create_user: CreateUser, authenticate_as: AuthenticateAs
) -> dict[str, str]:
    return authenticate_as(create_user(email="admin@example.com", role=Role.ADMIN))


@pytest.fixture
def mirante(create_condominium: CreateCondominium) -> Condominium:
    return create_condominium(name="Residencial Mirante")


@pytest.fixture
def aurora(create_condominium: CreateCondominium) -> Condominium:
    return create_condominium(name="Edifício Aurora")


@pytest.fixture
def ana(create_user: CreateUser) -> User:
    return create_user(email="ana@example.com", name="Ana", role=Role.OPERATOR)


@pytest.fixture
def bruno(create_user: CreateUser) -> User:
    return create_user(email="bruno@example.com", name="Bruno", role=Role.OPERATOR)


@pytest.fixture
def carla(create_user: CreateUser) -> User:
    """Operadora sem nenhum acordo."""
    return create_user(email="carla@example.com", name="Carla", role=Role.OPERATOR)


@pytest.fixture
def portfolio(
    db_session: Session,
    create_agreement_with_installments: CreateAgreementWithInstallments,
    create_debtor: CreateDebtor,
    mirante: Condominium,
    aurora: Condominium,
    ana: User,
    bruno: User,
    carla: User,
) -> None:
    """Parcelas de R$ 100,00 com R$ 10,00 de honorários cada.

    Fevereiro (mês fechado): vencem três parcelas da Ana. A de 10/02 foi paga só em
    01/03, então estava atrasada no fim do mês; a de 20/02 foi paga em 15/02; a de
    25/02 continua sem pagamento.

    Março (mês atual): pagas a de 01/03 do Bruno, no Aurora, e a de 01/03 do acordo
    sem responsável, além da de 10/02 da Ana, paga em 01/03. Os honorários
    recebidos em março somam R$ 30,00: Ana R$ 10,00, Bruno R$ 10,00 e R$ 10,00 sem
    responsável.
    """
    of_ana = create_agreement_with_installments(
        mirante,
        create_debtor(mirante, name="José"),
        [
            (date(2026, 2, 10), date(2026, 3, 1)),
            (date(2026, 2, 20), date(2026, 2, 15)),
            (date(2026, 2, 25), None),
        ],
    )
    of_ana.responsible_lawyer_id = ana.id
    of_bruno = create_agreement_with_installments(
        aurora,
        create_debtor(aurora, name="Maria"),
        [(date(2026, 3, 1), date(2026, 3, 1))],
    )
    of_bruno.responsible_lawyer_id = bruno.id
    create_agreement_with_installments(
        mirante,
        create_debtor(mirante, name="Paulo"),
        [(date(2026, 3, 1), date(2026, 3, 1))],
    )
    db_session.flush()


class TestAdminFinancialSummary:
    def test_office_revenue_and_each_members_commission(
        self,
        db_client: TestClient,
        admin_headers: dict[str, str],
        ana: User,
        bruno: User,
        carla: User,
        portfolio: None,
    ) -> None:
        """US33-CA01: inclui quem não tem comissão no mês, com zero."""
        response = db_client.get(FINANCIAL_URL, headers=admin_headers)

        assert response.status_code == 200
        assert response.json() == {
            "office_total_revenue": "30.00",
            "team_commissions": [
                {"user_id": str(ana.id), "user_name": "Ana", "commission": "10.00"},
                {"user_id": str(bruno.id), "user_name": "Bruno", "commission": "10.00"},
                {"user_id": str(carla.id), "user_name": "Carla", "commission": "0.00"},
            ],
            "period": {"month": 3, "year": 2026},
        }

    @pytest.mark.parametrize(
        "role",
        [
            pytest.param(Role.OPERATOR, id="operator"),
            pytest.param(Role.CLIENT, id="client"),
        ],
    )
    def test_only_the_admin_sees_it(
        self,
        db_client: TestClient,
        authenticate_as: AuthenticateAs,
        create_user: CreateUser,
        role: Role,
    ) -> None:
        """US33-CA02"""
        user = create_user(email="outro@example.com", role=role)

        response = db_client.get(FINANCIAL_URL, headers=authenticate_as(user))

        assert response.status_code == 403

    def test_requires_a_session(self, db_client: TestClient) -> None:
        assert db_client.get(FINANCIAL_URL).status_code == 401


class TestAdminDashboard:
    def test_indicators_of_a_closed_month(
        self,
        db_client: TestClient,
        admin_headers: dict[str, str],
        portfolio: None,
    ) -> None:
        """US34-CA01: em fevereiro venceram 3 parcelas e 2 estavam atrasadas no
        fim do mês (a de 10/02, paga só em março, e a de 25/02)."""
        response = db_client.get(
            DASHBOARD_URL, params={"year": 2026, "month": 2}, headers=admin_headers
        )

        assert response.status_code == 200
        assert response.json() == {
            "active_agreements": 3,
            "monthly_default_rate": 66.67,
            "overdue_installments": 2,
            "installments_due": 3,
            "paid_commissions": "10.00",
            "period": {"month": 2, "year": 2026},
        }

    def test_period_changes_the_indicators(
        self,
        db_client: TestClient,
        admin_headers: dict[str, str],
        portfolio: None,
    ) -> None:
        """US34-CA02"""
        march = db_client.get(
            DASHBOARD_URL, params={"year": 2026, "month": 3}, headers=admin_headers
        ).json()
        january = db_client.get(
            DASHBOARD_URL, params={"year": 2026, "month": 1}, headers=admin_headers
        ).json()

        assert march["installments_due"] == 2
        assert march["paid_commissions"] == "30.00"
        assert january["installments_due"] == 0
        assert january["monthly_default_rate"] == 0.0
        assert january["paid_commissions"] == "0.00"

    def test_current_month_counts_only_what_is_already_overdue_today(
        self,
        db_client: TestClient,
        db_session: Session,
        admin_headers: dict[str, str],
        create_agreement_with_installments: CreateAgreementWithInstallments,
        create_debtor: CreateDebtor,
        mirante: Condominium,
    ) -> None:
        """US34-CA01: hoje é 02/03; a parcela de 01/03 está atrasada, a de 10/03
        ainda não venceu."""
        create_agreement_with_installments(
            mirante,
            create_debtor(mirante),
            [(date(2026, 3, 1), None), (date(2026, 3, 10), None)],
        )
        db_session.flush()

        body = db_client.get(
            DASHBOARD_URL, params={"year": 2026, "month": 3}, headers=admin_headers
        ).json()

        assert body["installments_due"] == 2
        assert body["overdue_installments"] == 1
        assert body["monthly_default_rate"] == 50.0

    def test_future_month_has_nothing_overdue(
        self,
        db_client: TestClient,
        db_session: Session,
        admin_headers: dict[str, str],
        create_agreement_with_installments: CreateAgreementWithInstallments,
        create_debtor: CreateDebtor,
        mirante: Condominium,
    ) -> None:
        create_agreement_with_installments(
            mirante, create_debtor(mirante), [(date(2026, 4, 10), None)]
        )
        db_session.flush()

        body = db_client.get(
            DASHBOARD_URL, params={"year": 2026, "month": 4}, headers=admin_headers
        ).json()

        assert body["installments_due"] == 1
        assert body["overdue_installments"] == 0

    def test_cancelled_agreements_are_not_active(
        self,
        db_client: TestClient,
        db_session: Session,
        admin_headers: dict[str, str],
        create_agreement_with_installments: CreateAgreementWithInstallments,
        create_debtor: CreateDebtor,
        mirante: Condominium,
    ) -> None:
        """US34-CA01: o indicador segue os dados mais recentes."""
        agreement = create_agreement_with_installments(
            mirante, create_debtor(mirante), [(date(2026, 4, 10), None)]
        )
        agreement.cancel(None, datetime(2026, 3, 1, tzinfo=UTC))
        db_session.flush()

        body = db_client.get(
            DASHBOARD_URL, params={"year": 2026, "month": 3}, headers=admin_headers
        ).json()

        assert body["active_agreements"] == 0

    @pytest.mark.parametrize(
        "params",
        [
            pytest.param({"year": 2026, "month": 13}, id="month-13"),
            pytest.param({"year": 2026, "month": 0}, id="month-0"),
            pytest.param({"year": 2026}, id="no-month"),
        ],
    )
    def test_invalid_period_is_refused(
        self,
        db_client: TestClient,
        admin_headers: dict[str, str],
        params: dict[str, int],
    ) -> None:
        """US34-CA02"""
        response = db_client.get(DASHBOARD_URL, params=params, headers=admin_headers)

        assert response.status_code == 422

    def test_only_the_admin_sees_it(
        self,
        db_client: TestClient,
        authenticate_as: AuthenticateAs,
        ana: User,
    ) -> None:
        response = db_client.get(
            DASHBOARD_URL,
            params={"year": 2026, "month": 3},
            headers=authenticate_as(ana),
        )

        assert response.status_code == 403


def export_rows(
    client: TestClient, headers: dict[str, str], **params: Any
) -> list[tuple[Any, ...]]:
    response = client.get(EXPORT_URL, params=params, headers=headers)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    return list(workbook["Acompanhamento"].iter_rows(values_only=True))


class TestExportReport:
    def test_spreadsheet_has_the_installments_due_in_the_month(
        self,
        db_client: TestClient,
        admin_headers: dict[str, str],
        portfolio: None,
    ) -> None:
        """US35-CA01"""
        response = db_client.get(
            EXPORT_URL, params={"year": 2026, "month": 2}, headers=admin_headers
        )

        assert response.status_code == 200
        assert response.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert response.headers["content-disposition"] == (
            'attachment; filename="acompanhamento_2026_02.xlsx"'
        )
        rows = export_rows(db_client, admin_headers, year=2026, month=2)
        assert rows[0][0] == "Condomínio"
        assert [(row[1], row[4], row[7]) for row in rows[1:]] == [
            ("José", 1, "paid"),
            ("José", 2, "paid"),
            ("José", 3, "overdue"),
        ]

    @pytest.mark.parametrize(
        ("filters", "debtors"),
        [
            pytest.param({"status": "overdue"}, ["José"], id="status"),
            pytest.param({"status": "paid"}, ["Maria", "Paulo"], id="status-paid"),
        ],
    )
    def test_filters_by_status(
        self,
        db_client: TestClient,
        admin_headers: dict[str, str],
        portfolio: None,
        filters: dict[str, str],
        debtors: list[str],
    ) -> None:
        """US35-CA01"""
        month = 2 if filters["status"] == "overdue" else 3
        rows = export_rows(db_client, admin_headers, year=2026, month=month, **filters)

        assert sorted(row[1] for row in rows[1:]) == debtors

    def test_filters_open_installments(
        self,
        db_client: TestClient,
        db_session: Session,
        admin_headers: dict[str, str],
        create_agreement_with_installments: CreateAgreementWithInstallments,
        create_debtor: CreateDebtor,
        mirante: Condominium,
        portfolio: None,
    ) -> None:
        """US35-CA01: em aberto é a sem pagamento que ainda não venceu hoje."""
        create_agreement_with_installments(
            mirante, create_debtor(mirante, name="Rita"), [(date(2026, 3, 10), None)]
        )
        db_session.flush()

        rows = export_rows(db_client, admin_headers, year=2026, month=3, status="open")

        assert [(row[1], row[7]) for row in rows[1:]] == [("Rita", "open")]

    def test_filters_by_condominium(
        self,
        db_client: TestClient,
        admin_headers: dict[str, str],
        aurora: Condominium,
        portfolio: None,
    ) -> None:
        """US35-CA01"""
        rows = export_rows(
            db_client, admin_headers, year=2026, month=3, condominium_id=aurora.id
        )

        assert [(row[0], row[1]) for row in rows[1:]] == [("Edifício Aurora", "Maria")]

    def test_filters_by_responsible_lawyer(
        self,
        db_client: TestClient,
        admin_headers: dict[str, str],
        bruno: User,
        portfolio: None,
    ) -> None:
        rows = export_rows(
            db_client,
            admin_headers,
            year=2026,
            month=3,
            responsible_lawyer_id=bruno.id,
        )

        assert [(row[1], row[10]) for row in rows[1:]] == [("Maria", "Bruno")]

    def test_values_keep_their_format_in_excel(
        self,
        db_client: TestClient,
        admin_headers: dict[str, str],
        portfolio: None,
    ) -> None:
        """US35-CA02: valores são números e datas são datas, com formato próprio."""
        response = db_client.get(
            EXPORT_URL, params={"year": 2026, "month": 2}, headers=admin_headers
        )
        sheet = load_workbook(BytesIO(response.content))["Acompanhamento"]
        due, paid, amount, fee = (
            sheet.cell(row=2, column=col) for col in (6, 7, 9, 10)
        )

        assert isinstance(due.value, datetime)
        assert due.value.date() == date(2026, 2, 10)
        assert due.number_format == "dd/mm/yyyy"
        assert isinstance(paid.value, datetime)
        assert paid.value.date() == date(2026, 3, 1)
        assert amount.value == 100
        assert amount.number_format == "#,##0.00"
        assert fee.value == 10
        assert fee.number_format == "#,##0.00"

    def test_names_that_look_like_formulas_stay_text(
        self,
        db_client: TestClient,
        db_session: Session,
        admin_headers: dict[str, str],
        create_agreement_with_installments: CreateAgreementWithInstallments,
        create_debtor: CreateDebtor,
        mirante: Condominium,
    ) -> None:
        """US35-CA02: um nome digitado como fórmula não roda no Excel."""
        create_agreement_with_installments(
            mirante,
            create_debtor(mirante, name='=HYPERLINK("http://x")'),
            [(date(2026, 3, 10), None)],
        )
        db_session.flush()

        response = db_client.get(
            EXPORT_URL, params={"year": 2026, "month": 3}, headers=admin_headers
        )
        cell = load_workbook(BytesIO(response.content))["Acompanhamento"]["B2"]

        assert cell.data_type == "s"
        assert cell.value == '=HYPERLINK("http://x")'

    def test_month_without_installments_has_only_the_header(
        self, db_client: TestClient, admin_headers: dict[str, str]
    ) -> None:
        rows = export_rows(db_client, admin_headers, year=2026, month=1)

        assert len(rows) == 1

    def test_only_the_admin_exports(
        self,
        db_client: TestClient,
        authenticate_as: AuthenticateAs,
        ana: User,
    ) -> None:
        response = db_client.get(
            EXPORT_URL,
            params={"year": 2026, "month": 3},
            headers=authenticate_as(ana),
        )

        assert response.status_code == 403
